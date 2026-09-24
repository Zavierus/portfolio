import os
import sys
import time
import socket
import threading
from pathlib import Path
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer

from playwright.sync_api import sync_playwright
import pptx
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
DIST = ROOT / "dist-export"
SLIDES_DIR = DIST / "slides"


class SilentHandler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(ROOT), **kwargs)

    def log_message(self, format, *args):
        # Silence server logs
        pass


def get_free_port():
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


def start_server(port):
    server = ThreadingHTTPServer(("127.0.0.1", port), SilentHandler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    return server


SLIDE_NOTES = [
    "P01: 封面 - ZIAVER / 王泽源 · 内容产品与创意技术 (2026 Edition)",
    "P02: 理念与工作流 - 三位一体工作流体系 (内容策略 × 视觉审美 × 可玩原型)",
    "P03: 核心战绩 - 业务结果数据看板 (25.81亿 GMV / 4.7% CTR / 2000+ 摄影人次 / 56万 粉丝)",
    "P04: 职业经历 - 字节跳动、瓜子二手车、敢览文化实战历程与能力雷达",
    "P05: 互动原型 01 - FRAME//ZERO (3D 时间控制第一人称射击战斗切片)",
    "P06: 互动原型 02 - PULSE ROOM (基于 Mineradio 2.0.2 的雾气播放器与响应式声场)",
    "P07: 互动原型 03 - SET//FLOW (直播空间规划与机位视口 3D 工作台)",
    "P08: 互动原型 04 - PLAYER SIGNAL (Steam 真实玩家反馈与情报分析看板)",
    "P09: 视觉展陈 - RELIC//01 (生物机械材质分层与实时着色器实验展陈)",
    "P10: 视觉档案 - 商业摄影接触表精选 & 图像平面系统 & 微信动态推文",
    "P11: 封底与联系 - 联系方式与全站在线可交互项目扫码体验",
]


def get_writable_path(path: Path) -> Path:
    try:
        if path.exists():
            with open(path, "a+b"):
                pass
        return path
    except (PermissionError, OSError):
        alt = path.with_name(path.stem + "_最新版" + path.suffix)
        print(f"  [NOTE] {path.name} is currently open in another program. Writing to {alt.name}")
        return alt


def build_pptx(slide_images, output_path):
    output_path = get_writable_path(output_path)
    prs = pptx.Presentation()
    # 16:9 widescreen: 13.333 x 7.5 inches
    prs.slide_width = Inches(13.3333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]  # blank layout

    for idx, img_path in enumerate(slide_images):
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(str(img_path), Inches(0), Inches(0), width=prs.slide_width, height=prs.slide_height)
        notes_slide = slide.notes_slide
        text_frame = notes_slide.notes_text_frame
        note_text = SLIDE_NOTES[idx] if idx < len(SLIDE_NOTES) else f"Slide {idx + 1}"
        text_frame.text = note_text

    prs.save(str(output_path))
    print(f"  [OK] PPTX generated: {output_path.name} ({len(slide_images)} slides)")


def main():
    DIST.mkdir(parents=True, exist_ok=True)
    SLIDES_DIR.mkdir(parents=True, exist_ok=True)

    port = get_free_port()
    server = start_server(port)
    print(f"[+] Local HTTP server started at http://127.0.0.1:{port}")

    edge_paths = [
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
    ]
    executable = next((p for p in edge_paths if os.path.exists(p)), None)
    if executable:
        print(f"[+] Using Edge browser: {executable}")
    else:
        print("[+] Using Playwright default chromium")

    with sync_playwright() as p:
        launch_args = {
            "headless": True,
            "args": ["--enable-webgl", "--ignore-gpu-blocklist", "--font-render-hinting=none"],
        }
        if executable:
            launch_args["executable_path"] = executable

        browser = p.chromium.launch(**launch_args)

        # -------------------------------------------------------------
        # 1. Export 16:9 Deck PDF & PPTX
        # -------------------------------------------------------------
        print("\n[1/4] Generating 16:9 Presentation PDF & PPTX...")
        deck_ctx = browser.new_context(
            viewport={"width": 1920, "height": 1080},
            device_scale_factor=2,  # Retina sharp
        )
        deck_page = deck_ctx.new_page()
        deck_url = f"http://127.0.0.1:{port}/deck/index.html"
        deck_page.goto(deck_url, wait_until="networkidle")
        deck_page.wait_for_timeout(1000)

        # 1a. Export Multi-page 16:9 PDF
        deck_pdf_path = get_writable_path(DIST / "ZIAVER_王泽源_作品集画册_16x9.pdf")
        deck_page.pdf(
            path=str(deck_pdf_path),
            width="1920px",
            height="1080px",
            print_background=True,
            prefer_css_page_size=True,
            margin={"top": "0", "bottom": "0", "left": "0", "right": "0"},
        )
        print(f"  [OK] PDF generated: {deck_pdf_path.name}")

        # 1b. Capture each slide as high-res PNG for PPTX
        slide_elements = deck_page.locator(".slide").all()
        slide_imgs = []
        for i, slide in enumerate(slide_elements):
            img_path = SLIDES_DIR / f"slide_{i+1:02d}.png"
            slide.screenshot(path=str(img_path))
            slide_imgs.append(img_path)
        print(f"  [OK] Exported {len(slide_imgs)} high-resolution slide PNGs")

        # 1c. Build PPTX
        deck_pptx_path = get_writable_path(DIST / "ZIAVER_王泽源_作品集画册_16x9.pptx")
        build_pptx(slide_imgs, deck_pptx_path)
        deck_ctx.close()

        # -------------------------------------------------------------
        # 2. Export Resume Showcase Card (简历专用一图流展板)
        # -------------------------------------------------------------
        print("\n[2/4] Generating Resume Showcase Card (简历专用一图流展板)...")
        card_ctx = browser.new_context(
            viewport={"width": 2600, "height": 1300},
            device_scale_factor=2,  # Ultra-crisp 4800px wide
        )
        card_page = card_ctx.new_page()
        card_url = f"http://127.0.0.1:{port}/deck/resume-card.html"
        card_page.goto(card_url, wait_until="networkidle")
        card_page.wait_for_timeout(1000)

        card_locator = card_page.locator(".showcase-card")
        card_png_path = get_writable_path(DIST / "ZIAVER_简历专用_作品集一图流展板.png")
        card_locator.screenshot(path=str(card_png_path))
        print(f"  [OK] Resume card generated: {card_png_path.name}")
        card_ctx.close()

        # -------------------------------------------------------------
        # 3. Export Full-Page Long Scroll Showcase (全景超清长图一图流)
        # -------------------------------------------------------------
        print("\n[3/4] Generating Full-Length Long Scroll Showcase (全景超清长图)...")
        long_ctx = browser.new_context(
            viewport={"width": 1440, "height": 900},
            device_scale_factor=2,  # 2880px crisp width
        )
        long_page = long_ctx.new_page()
        home_url = f"http://127.0.0.1:{port}/index.html"
        long_page.goto(home_url, wait_until="networkidle")
        long_page.wait_for_timeout(1000)

        # Force reveal elements, absolute header, and guarantee portal door preview images
        long_page.evaluate("""() => {
            document.querySelectorAll('.reveal').forEach(el => el.classList.add('is-visible'));
            const header = document.querySelector('.site-header');
            if (header) header.style.position = 'absolute';

            const style = document.createElement('style');
            style.textContent = `
                .portal-door > img {
                    opacity: 1 !important;
                    visibility: visible !important;
                    display: block !important;
                }
                .portal-door > .project-media__canvas {
                    display: none !important;
                }
                .portal.is-project-preview-live .portal-door > img {
                    opacity: 1 !important;
                }
            `;
            document.head.appendChild(style);
        }""")
        
        # Scroll down progressively to mount canvases and lazy assets
        scroll_height = long_page.evaluate("() => document.body.scrollHeight")
        for pos in range(0, scroll_height, 500):
            long_page.evaluate(f"window.scrollTo(0, {pos})")
            long_page.wait_for_timeout(120)
        long_page.evaluate("window.scrollTo(0, 0)")
        long_page.wait_for_timeout(1200)

        long_png_path = get_writable_path(DIST / "ZIAVER_作品集全景长图_一图流.png")
        long_page.screenshot(path=str(long_png_path), full_page=True)
        print(f"  [OK] Long scroll generated: {long_png_path.name}")
        long_ctx.close()

        # -------------------------------------------------------------
        # 4. Export Webpage Original Print PDF (网页原貌打印版)
        # -------------------------------------------------------------
        print("\n[4/4] Generating Webpage Print PDF (网页原貌打印版)...")
        print_ctx = browser.new_context(
            viewport={"width": 1200, "height": 1600},
            device_scale_factor=2,
        )
        print_page = print_ctx.new_page()
        print_page.goto(home_url, wait_until="networkidle")
        print_page.wait_for_timeout(1000)

        # Force reveal elements, force eager images and scroll through
        print_page.evaluate("""() => {
            document.querySelectorAll('.reveal').forEach(el => el.classList.add('is-visible'));
            document.querySelectorAll('img').forEach(img => {
                img.loading = 'eager';
                img.removeAttribute('loading');
                img.decoding = 'sync';
            });
            const header = document.querySelector('.site-header');
            if (header) header.style.position = 'absolute';
        }""")
        p_height = print_page.evaluate("() => document.body.scrollHeight")
        for pos in range(0, p_height, 400):
            print_page.evaluate(f"window.scrollTo(0, {pos})")
            print_page.wait_for_timeout(80)
        print_page.evaluate("window.scrollTo(0, 0)")
        print_page.wait_for_timeout(800)

        # Ensure all images are loaded before printing
        print_page.evaluate("""() => Promise.all(
            Array.from(document.images).map(img => {
                if (img.complete) return Promise.resolve();
                return new Promise(resolve => {
                    img.onload = img.onerror = resolve;
                });
            })
        )""")

        # Inject print CSS optimizations (严格 6 页黄金排版)
        print_page.add_style_tag(content="""
            @media print {
                @page {
                    size: A4 portrait;
                    margin: 8mm 10mm;
                }
                * {
                    -webkit-print-color-adjust: exact !important;
                    print-color-adjust: exact !important;
                    box-sizing: border-box !important;
                }
                html, body {
                    background: #070808 !important;
                    color: #ebeae4 !important;
                    font-size: 12px !important;
                    margin: 0 !important;
                    padding: 0 !important;
                }
                .site-header, .skip-link, .language-toggle, .hr-search,
                [data-atmosphere-canvas], .project-transition, .gallery-lightbox,
                .identity-cue, .portal-gallery__status, .visual-archive__all,
                .visual-editor__bar a, body::before {
                    display: none !important;
                }
                .reveal {
                    opacity: 1 !important;
                    transform: none !important;
                    visibility: visible !important;
                }

                section {
                    min-height: auto !important;
                    height: auto !important;
                    padding: 0 !important;
                    margin: 0 !important;
                }

                /* PAGE 1: Identity / Hero (全屏画册级震撼封面) */
                .identity {
                    height: 980px !important;
                    max-height: 980px !important;
                    min-height: 980px !important;
                    position: relative !important;
                    overflow: hidden !important;
                    page-break-after: always !important;
                    break-after: page !important;
                    background: #070808 !important;
                }
                .identity-visual {
                    position: absolute !important;
                    inset: 0 !important;
                    z-index: 0 !important;
                    display: block !important;
                    overflow: hidden !important;
                }
                .identity-visual canvas {
                    display: none !important;
                }
                .identity-visual img {
                    position: absolute !important;
                    inset: 0 !important;
                    width: 100% !important;
                    height: 100% !important;
                    object-fit: cover !important;
                    object-position: center 48% !important;
                    display: block !important;
                    z-index: 1 !important;
                    opacity: 1 !important;
                    visibility: visible !important;
                    filter: brightness(0.85) contrast(1.1) !important;
                }
                .identity-scrim {
                    position: absolute !important;
                    inset: 0 !important;
                    z-index: 2 !important;
                    display: block !important;
                    width: 100% !important;
                    background: linear-gradient(90deg, rgba(7,8,8,0.92) 0%, rgba(7,8,8,0.65) 45%, rgba(7,8,8,0.15) 100%) !important;
                    border: none !important;
                }
                .identity-copy {
                    position: absolute !important;
                    bottom: 80px !important;
                    left: 24px !important;
                    z-index: 3 !important;
                    max-width: 600px !important;
                    padding: 0 !important;
                }
                .identity-copy .utility-label {
                    font-family: monospace !important;
                    font-size: 11px !important;
                    color: #8effca !important;
                    margin-bottom: 6px !important;
                    letter-spacing: 0.1em !important;
                }
                .identity-copy h1 {
                    font-size: 56px !important;
                    line-height: 0.95 !important;
                    margin: 4px 0 8px !important;
                    color: #ffffff !important;
                }
                .identity-copy h1 small {
                    font-size: 28px !important;
                    margin-left: 12px !important;
                    color: #ebeae4 !important;
                }
                .identity-role {
                    font-size: 20px !important;
                    color: #d9ff32 !important;
                    margin-bottom: 8px !important;
                    font-weight: 700 !important;
                }
                .identity-intro {
                    font-size: 14px !important;
                    color: #cbd5e1 !important;
                    line-height: 1.6 !important;
                    margin: 0 !important;
                }
                .identity-actions {
                    margin-top: 16px !important;
                    display: flex !important;
                    gap: 12px !important;
                }
                .identity-actions a {
                    border: 1px solid rgba(255, 255, 255, 0.25) !important;
                    border-radius: 4px !important;
                    padding: 9px 20px !important;
                    font-size: 11px !important;
                    color: #fff !important;
                    text-decoration: none !important;
                }
                .identity-actions a:first-child {
                    background: #d9ff32 !important;
                    color: #000 !important;
                    font-weight: 700 !important;
                }
                .identity-project {
                    position: absolute !important;
                    bottom: 40px !important;
                    right: 24px !important;
                    z-index: 3 !important;
                    font-family: monospace !important;
                    font-size: 11px !important;
                    color: #94a3b8 !important;
                }

                /* PAGE 2: Selected Works / 4 Portals (精确 1 页) */
                .work {
                    height: 980px !important;
                    max-height: 980px !important;
                    padding: 16px 0 !important;
                    page-break-after: always !important;
                    break-after: page !important;
                    display: flex !important;
                    flex-direction: column !important;
                    justify-content: space-between !important;
                }
                .portal-gallery {
                    display: flex !important;
                    flex-direction: column !important;
                    height: 100% !important;
                    justify-content: space-between !important;
                    background: none !important;
                    padding: 0 !important;
                }
                .portal-gallery__head {
                    margin-bottom: 12px !important;
                }
                .portal-gallery__head h2 {
                    font-size: 32px !important;
                    line-height: 1.05 !important;
                    margin: 4px 0 !important;
                    color: #ffffff !important;
                }
                .portal-gallery__head h2 em {
                    color: #d9ff32 !important;
                    font-style: normal !important;
                }
                .portal-gallery__head p {
                    font-size: 12px !important;
                    color: #a6aaa7 !important;
                    margin: 0 !important;
                }
                .portal-gallery__field {
                    display: grid !important;
                    grid-template-columns: 1fr 1fr !important;
                    grid-template-rows: 390px 390px !important;
                    gap: 16px !important;
                    margin: 0 !important;
                    perspective: none !important;
                }
                .portal {
                    opacity: 1 !important;
                    transform: none !important;
                    filter: none !important;
                    height: 100% !important;
                    display: block !important;
                    text-decoration: none !important;
                }
                .portal-door {
                    height: 100% !important;
                    border: 1px solid rgba(255, 255, 255, 0.16) !important;
                    border-radius: 8px !important;
                    overflow: hidden !important;
                    position: relative !important;
                    display: block !important;
                    background: #0d1214 !important;
                }
                .portal-door > img {
                    opacity: 1 !important;
                    visibility: visible !important;
                    display: block !important;
                    width: 100% !important;
                    height: 100% !important;
                    object-fit: cover !important;
                }
                .portal-door > .project-media__canvas {
                    display: none !important;
                }
                .portal-door__meta {
                    opacity: 1 !important;
                    background: linear-gradient(0deg, rgba(7, 8, 8, 0.95) 0%, rgba(7, 8, 8, 0.65) 60%, transparent 100%) !important;
                    padding: 16px 20px !important;
                    position: absolute !important;
                    bottom: 0 !important;
                    left: 0 !important;
                    right: 0 !important;
                }
                .portal-door__meta small {
                    font-family: monospace !important;
                    font-size: 10px !important;
                    letter-spacing: 0.12em !important;
                    color: #8effca !important;
                    display: block !important;
                }
                .portal-door__meta b {
                    font-size: 20px !important;
                    display: block !important;
                    margin: 3px 0 !important;
                    color: #ffffff !important;
                }
                .portal-door__meta em {
                    font-size: 12px !important;
                    color: #d4e0d9 !important;
                    display: block !important;
                    font-style: normal !important;
                }
                .portal-door__action {
                    display: none !important;
                }

                /* PAGE 3: Outcomes (权威居中展陈) */
                .outcomes {
                    height: 980px !important;
                    max-height: 980px !important;
                    padding: 24px 0 !important;
                    page-break-after: always !important;
                    break-after: page !important;
                    background: none !important;
                    display: flex !important;
                    flex-direction: column !important;
                    justify-content: center !important;
                }
                .outcomes .section-heading {
                    padding: 0 0 24px 0 !important;
                    margin: 0 !important;
                }
                .outcomes .section-heading h2 {
                    font-size: 36px !important;
                    line-height: 1.05 !important;
                    margin: 6px 0 !important;
                    color: #ffffff !important;
                }
                .outcomes .section-heading p {
                    font-size: 13px !important;
                    color: #a5b8ae !important;
                }
                .outcome-ledger {
                    display: grid !important;
                    grid-template-columns: 1fr 1fr !important;
                    gap: 20px !important;
                    margin: 0 !important;
                    border: none !important;
                }
                .outcome-ledger article {
                    min-height: 220px !important;
                    border: 1px solid rgba(213, 242, 231, 0.22) !important;
                    border-radius: 8px !important;
                    padding: 28px 30px !important;
                    background: rgba(13, 23, 22, 0.95) !important;
                }
                .outcome-ledger strong {
                    font-size: 44px !important;
                    -webkit-text-fill-color: #ff5a36 !important;
                    color: #ff5a36 !important;
                    line-height: 1 !important;
                    display: block !important;
                }
                .outcome-ledger h3 {
                    margin-top: 14px !important;
                    font-size: 17px !important;
                    color: #ffffff !important;
                }
                .outcome-ledger p {
                    margin-top: 8px !important;
                    font-size: 13px !important;
                    line-height: 1.6 !important;
                    color: #a5b8ae !important;
                }

                /* PAGE 4: Experience (经历 + 能力索引，严格 1 页) */
                .experience {
                    height: 980px !important;
                    max-height: 980px !important;
                    padding: 20px 0 !important;
                    page-break-after: always !important;
                    break-after: page !important;
                    background: none !important;
                    display: flex !important;
                    flex-direction: column !important;
                    justify-content: space-between !important;
                }
                .experience .section-heading {
                    padding: 0 !important;
                    margin: 0 0 10px 0 !important;
                }
                .experience .section-heading h2 {
                    font-size: 26px !important;
                    line-height: 1.1 !important;
                    margin: 4px 0 !important;
                    color: #ffffff !important;
                }
                .experience .section-heading p {
                    font-size: 11px !important;
                    color: #8c9ba0 !important;
                    margin: 0 !important;
                }
                .experience-index {
                    margin: 0 !important;
                    border: none !important;
                    display: flex !important;
                    flex-direction: column !important;
                    gap: 8px !important;
                }
                .experience-item {
                    background: #0d1214 !important;
                    border: 1px solid rgba(255, 255, 255, 0.1) !important;
                    border-radius: 6px !important;
                    padding: 10px 18px !important;
                    break-inside: avoid !important;
                    page-break-inside: avoid !important;
                }
                .experience-item__top {
                    display: flex !important;
                    justify-content: space-between !important;
                    align-items: baseline !important;
                    margin-bottom: 4px !important;
                }
                .experience-item h3 {
                    font-size: 14px !important;
                    color: #ffffff !important;
                    margin: 0 0 2px 0 !important;
                }
                .experience-item .company-role {
                    font-size: 11px !important;
                    color: #d9ff32 !important;
                }
                .experience-item time {
                    font-family: monospace !important;
                    font-size: 10px !important;
                    color: #7b8b91 !important;
                }
                .experience-signals {
                    display: flex !important;
                    gap: 24px !important;
                    margin-top: 4px !important;
                    border-top: 1px solid rgba(255, 255, 255, 0.06) !important;
                    padding-top: 4px !important;
                }
                .experience-signals span strong {
                    font-size: 14px !important;
                    color: #ffffff !important;
                    display: block !important;
                }
                .experience-signals span small {
                    font-size: 9px !important;
                    color: #7b8b91 !important;
                }
                .capability-index {
                    margin: 8px 0 0 !important;
                    background: #0d1214 !important;
                    border: 1px solid rgba(255, 255, 255, 0.1) !important;
                    border-radius: 6px !important;
                    padding: 10px 18px !important;
                    display: flex !important;
                    justify-content: space-between !important;
                    gap: 16px !important;
                    break-inside: avoid !important;
                    page-break-inside: avoid !important;
                }
                .capability-index > div {
                    flex: 1 !important;
                }
                .capability-index span {
                    font-family: monospace !important;
                    font-size: 9px !important;
                    color: #d9ff32 !important;
                    display: block !important;
                    margin-bottom: 2px !important;
                    font-weight: 700 !important;
                }
                .capability-index p {
                    font-size: 10px !important;
                    color: #a2b1b7 !important;
                    margin: 0 !important;
                    line-height: 1.4 !important;
                }

                /* PAGE 5: Visual Archive (摄影接触表 + 平面推文，严格 1 页) */
                .visual-archive {
                    min-height: auto !important;
                    padding: 16px 0 0 0 !important;
                    background: none !important;
                }
                .visual-archive__heading {
                    margin-bottom: 10px !important;
                }
                .visual-archive__heading h2 {
                    font-size: 28px !important;
                    line-height: 1 !important;
                    margin: 2px 0 4px 0 !important;
                    color: #ffffff !important;
                }
                .visual-archive__heading p {
                    font-size: 11px !important;
                    color: #8c9ba0 !important;
                    margin: 0 !important;
                }
                .visual-editor {
                    display: flex !important;
                    flex-direction: column !important;
                    justify-content: space-between !important;
                    page-break-after: always !important;
                    break-after: page !important;
                }
                .visual-editor__bar {
                    margin-bottom: 8px !important;
                }
                .visual-editor__bar span {
                    font-family: monospace !important;
                    font-size: 10px !important;
                    color: #8effca !important;
                }
                .visual-editor__bar h3 {
                    font-size: 15px !important;
                    color: #ffffff !important;
                    margin: 2px 0 0 0 !important;
                }
                .contact-sheet-layout {
                    display: block !important;
                }
                .contact-detail {
                    display: none !important;
                }
                .contact-sheet {
                    display: grid !important;
                    grid-template-columns: repeat(6, 1fr) !important;
                    gap: 6px !important;
                    padding: 0 !important;
                    margin: 0 0 10px 0 !important;
                    list-style: none !important;
                }
                .contact-sheet li, .contact-sheet button {
                    display: block !important;
                    border: none !important;
                    background: none !important;
                    padding: 0 !important;
                    margin: 0 !important;
                }
                .contact-sheet img {
                    width: 100% !important;
                    height: 82px !important;
                    object-fit: cover !important;
                    border-radius: 4px !important;
                    display: block !important;
                }
                .visual-editor__companions {
                    display: grid !important;
                    grid-template-columns: 1fr 1fr !important;
                    gap: 12px !important;
                    margin-top: 8px !important;
                }
                .visual-archive__lane {
                    background: #0d1214 !important;
                    border: 1px solid rgba(255, 255, 255, 0.09) !important;
                    border-radius: 8px !important;
                    overflow: hidden !important;
                    display: flex !important;
                    flex-direction: row !important;
                    height: 100px !important;
                    text-decoration: none !important;
                }
                .visual-archive__image {
                    width: 100px !important;
                    height: 100px !important;
                    flex-shrink: 0 !important;
                }
                .visual-archive__image img {
                    width: 100% !important;
                    height: 100% !important;
                    object-fit: cover !important;
                }
                .visual-archive__meta {
                    padding: 10px 14px !important;
                    display: flex !important;
                    flex-direction: column !important;
                    justify-content: center !important;
                }
                .visual-archive__meta small {
                    font-family: monospace !important;
                    font-size: 9px !important;
                    color: #d9ff32 !important;
                }
                .visual-archive__meta strong {
                    font-size: 13px !important;
                    color: #fff !important;
                    margin: 2px 0 !important;
                }
                .visual-archive__meta em {
                    font-size: 10px !important;
                    color: #8c9ba0 !important;
                    font-style: normal !important;
                }
                .visual-archive__meta b {
                    font-size: 9px !important;
                    color: #8effca !important;
                    margin-top: 3px !important;
                }

                /* PAGE 6: Relic Exhibit + Contact (严格第 6 页) */
                .relic-exhibit {
                    background: #0d1214 !important;
                    border: 1px solid rgba(255, 90, 54, 0.28) !important;
                    border-radius: 8px !important;
                    overflow: hidden !important;
                    display: grid !important;
                    grid-template-columns: 1.1fr 1fr !important;
                    margin-top: 24px !important;
                    margin-bottom: 24px !important;
                    break-inside: avoid !important;
                }
                .relic-exhibit__image img {
                    width: 100% !important;
                    height: 280px !important;
                    object-fit: cover !important;
                    display: block !important;
                }
                .relic-exhibit__copy {
                    padding: 22px 26px !important;
                    display: flex !important;
                    flex-direction: column !important;
                    justify-content: center !important;
                }
                .relic-exhibit__copy small {
                    font-family: monospace !important;
                    font-size: 10px !important;
                    color: #ff5a36 !important;
                }
                .relic-exhibit__copy strong {
                    font-size: 19px !important;
                    color: #fff !important;
                    margin: 6px 0 !important;
                }
                .relic-exhibit__copy em {
                    font-size: 12px !important;
                    color: #8c9ba0 !important;
                    line-height: 1.6 !important;
                    font-style: normal !important;
                }
                .relic-exhibit__copy b {
                    font-size: 10px !important;
                    color: #ff5a36 !important;
                    margin-top: 8px !important;
                }

                .contact {
                    min-height: auto !important;
                    padding: 20px 0 16px !important;
                    border-top: 1px solid rgba(255, 255, 255, 0.1) !important;
                    break-inside: avoid !important;
                }
                .contact .utility-label {
                    font-family: monospace !important;
                    font-size: 10px !important;
                    color: #8effca !important;
                    margin-bottom: 4px !important;
                }
                .contact h2 {
                    font-size: 28px !important;
                    line-height: 1.1 !important;
                    margin: 4px 0 8px !important;
                    color: #ffffff !important;
                }
                .contact p {
                    font-size: 12px !important;
                    color: #8c9ba0 !important;
                    margin-bottom: 14px !important;
                }
                .contact-actions {
                    display: flex !important;
                    gap: 12px !important;
                }
                .contact-actions a {
                    border: 1px solid rgba(255, 255, 255, 0.18) !important;
                    border-radius: 4px !important;
                    padding: 8px 16px !important;
                    font-size: 11px !important;
                    color: #fff !important;
                    text-decoration: none !important;
                }
                .contact-actions a:first-child {
                    background: #d9ff32 !important;
                    color: #000 !important;
                    font-weight: 700 !important;
                }

                .site-footer {
                    padding: 16px 0 !important;
                    border-top: 1px solid rgba(255, 255, 255, 0.08) !important;
                    display: flex !important;
                    justify-content: space-between !important;
                    font-family: monospace !important;
                    font-size: 10px !important;
                    color: #617075 !important;
                }
            }
        """)
        long_pdf_path = get_writable_path(DIST / "ZIAVER_王泽源_作品集网页原貌_打印版.pdf")
        print_page.pdf(
            path=str(long_pdf_path),
            format="A4",
            print_background=True,
            margin={"top": "8mm", "bottom": "8mm", "left": "10mm", "right": "10mm"},
        )
        print(f"  [OK] Print PDF generated: {long_pdf_path.name}")
        print_ctx.close()

        browser.close()

    server.shutdown()
    print("\n[+] All exports finished successfully!")
    print(f"[+] Output directory: {DIST.resolve()}")
    for item in sorted(DIST.iterdir()):
        if item.is_file():
            size_mb = item.stat().st_size / (1024 * 1024)
            print(f"    - {item.name:<45} ({size_mb:.2f} MB)")


if __name__ == "__main__":
    main()
